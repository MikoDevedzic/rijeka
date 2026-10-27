"""Excel replications: workbooks build, carry the formulas and the engine reference, products list is sane."""
import io
from datetime import date

import pytest
from openpyxl import load_workbook

from pricing.curve import Curve
from methodology import replicate as rep


def _mkt(flat=0.04):
    val = date(2026, 8, 3)
    # pillars every year to 30y on a flat 4% continuously compounded curve
    import math
    pil = [(val, 1.0)] + [(date(val.year + k, val.month, val.day), math.exp(-flat * k)) for k in range(1, 31)]
    pil.insert(1, (date(2026, 11, 3), math.exp(-flat * 0.25)))
    curve = Curve(val, df_pillars=pil)
    vol = [{"tenor": "1Yx5Y", "rate": 90.0, "enabled": True}]
    return rep.Market(valuation_date=val, curve_id="USD_SOFR", snapshot_date=val, source="TEST",
                      quotes=[{"tenor": "5Y", "quote_type": "OIS", "rate": 4.0, "enabled": True}],
                      curve=curve, vol_quotes=vol, vol_date=val)


@pytest.mark.parametrize("product", list(rep.BUILDERS))
def test_workbook_builds_with_formulas_and_reference(product):
    data, ref = rep.build(product, _mkt())
    wb = load_workbook(io.BytesIO(data))
    assert "README" in wb.sheetnames and "Curve" in wb.sheetnames
    assert "PillarDates" in wb.defined_names and "PillarDF" in wb.defined_names and "ValDate" in wb.defined_names
    formulas = [c.value for ws in wb.worksheets for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert len(formulas) > 10
    assert any("PillarDF" in f for f in formulas), "DF interpolation must reference the curve pillars"
    if product in ("swaption", "cap", "crypto_ndo"):
        assert any("NORM.S.DIST" in f for f in formulas)
    assert isinstance(ref, float)


def test_df_formula_matches_curve_interpolation():
    """The Excel DF formula is log-linear between bracketing pillar dates, like pricing/curve.py."""
    f = rep.DF("A9")
    assert "MATCH(A9,PillarDates,1)" in f and "LN(" in f and "EXP(" in f and "ROWS(PillarDates)-1" in f


def test_irs_reference_is_the_engine_value():
    """The green reference numbers are Rijeka's engine output, not re-derived in the sheet."""
    mkt = _mkt()
    data, npv = rep.build("irs", mkt)
    wb = load_workbook(io.BytesIO(data))
    assert wb["Summary"]["B2"].value == pytest.approx(npv)
    # pay 4% on a flat 4% (cont.) curve: par is a bit above 4% simple, so the payer swap has positive value
    assert 0 < npv < 300_000


def test_products_list_covers_crypto_and_post_trade():
    ids = {p["id"] for p in rep.PRODUCTS}
    assert {"irs", "ois", "swaption", "cap", "crypto_ndf", "crypto_ndo", "xva", "simm", "confirm"} <= ids
    for p in rep.PRODUCTS:
        assert p["replication"] in ("xlsx", "doc", "xlsx-static", "verifier")
        if p["replication"] == "xlsx":
            assert p["id"] in rep.BUILDERS


def test_route_rejects_unknown_product(monkeypatch):
    from fastapi.testclient import TestClient
    from unittest.mock import MagicMock
    from main import app
    from db.session import get_db
    from middleware.auth import verify_token
    app.dependency_overrides[get_db] = lambda: MagicMock()
    app.dependency_overrides[verify_token] = lambda: {"sub": "00000000-0000-0000-0000-000000000000"}
    try:
        c = TestClient(app)
        assert c.get("/api/methodology/nope/replication.xlsx").status_code == 404
        r = c.get("/api/methodology/products"); assert r.status_code == 200 and len(r.json()) >= 9
    finally:
        app.dependency_overrides.pop(get_db, None); app.dependency_overrides.pop(verify_token, None)


def test_sheet_arithmetic_reproduces_engine_legs():
    """
    No spreadsheet engine here, so evaluate the leg formulas' arithmetic in
    Python with the same inputs the sheet gets (schedule dates as values, DFs
    by the curve's own interpolation, ACT/360 and 30/360 day counts) and
    compare with Rijeka's leg PVs. This is the parity the green cells claim.
    """
    from datetime import date
    from pricing.day_count import thirty360, act360
    from pricing.schedule import generate_schedule
    from pricing.ir_swap import price_swap
    from decimal import Decimal
    mkt = _mkt(); val = mkt.valuation_date; spot = rep._spot(val); mat = date(spot.year + 5, spot.month, spot.day)
    N, r = 10_000_000, 0.04
    legs = [
        {"id": "f", "leg_ref": "FIXED", "leg_type": "FIXED", "direction": "PAY", "currency": "USD", "notional": N, "day_count": "30/360",
         "payment_frequency": "SEMI_ANNUAL", "bdc": "MOD_FOLLOWING", "fixed_rate": r, "effective_date": spot.isoformat(), "maturity_date": mat.isoformat(),
         "discount_curve_id": "USD_SOFR", "forecast_curve_id": "USD_SOFR"},
        {"id": "l", "leg_ref": "FLOAT", "leg_type": "FLOAT", "direction": "RECEIVE", "currency": "USD", "notional": N, "day_count": "ACT/360",
         "payment_frequency": "QUARTERLY", "bdc": "MOD_FOLLOWING", "spread": 0.0, "effective_date": spot.isoformat(), "maturity_date": mat.isoformat(),
         "discount_curve_id": "USD_SOFR", "forecast_curve_id": "USD_SOFR"},
    ]
    res = price_swap("t", legs, {"USD_SOFR": mkt.curve}, val)
    eng = {l.leg_ref: abs(l.pv) for l in res.legs}
    def leg(freq, dc, is_float, df_at):
        tot = 0.0
        for p in generate_schedule(spot, mat, freq, dc, Decimal(N), "MOD_FOLLOWING", is_float=is_float):
            tau = float((act360 if dc == "ACT/360" else thirty360)(p.period_start, p.period_end))
            rate = (mkt.curve.df(p.period_start) / mkt.curve.df(p.period_end) - 1) / tau if is_float else r
            tot += N * rate * tau * mkt.curve.df(getattr(p, df_at))
        return tot
    # Engine and sheet both discount at the business-day-adjusted payment date
    # (the date the cash moves). Same arithmetic, machine precision:
    assert leg("SEMI_ANNUAL", "30/360", False, "payment_date") == pytest.approx(eng["FIXED"], rel=1e-9)
    assert leg("QUARTERLY", "ACT/360", True, "payment_date") == pytest.approx(eng["FLOAT"], rel=1e-9)
    # Discounting at the unadjusted period end (the pre-v0.5 behaviour) differs
    # only where a date rolled over a weekend: a day or two, a few $ per million.
    assert leg("SEMI_ANNUAL", "30/360", False, "period_end") == pytest.approx(eng["FIXED"], rel=2e-4)
    assert leg("SEMI_ANNUAL", "30/360", False, "period_end") != pytest.approx(eng["FIXED"], rel=1e-9)
